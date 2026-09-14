from openpyxl import load_workbook
import csv

wb = load_workbook("radio_tracks.xlsx")
ws = wb["Tracks"]

with open("radio_tracks.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    for row in ws.iter_rows(values_only=True):
        writer.writerow([str(v) if v is not None else "" for v in row])